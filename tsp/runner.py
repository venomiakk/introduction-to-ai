from __future__ import annotations

import json
import multiprocessing
import os
import pickle
import re
import tempfile
from pathlib import Path
from queue import Empty
from typing import Any, Callable

from openpyxl import Workbook, load_workbook

from algorithms.a_star import run_astar
from algorithms.aco import run_aco
from algorithms.bfs import run_bfs
from algorithms.dfs import run_dfs
from algorithms.dijkstra import run_dijkstra
from algorithms.nn import run_nn
from measure_perf import measure_performance


DATA_DIRECTORY = Path(__file__).resolve().parent / "data"
OUTPUT_FILE = Path(__file__).resolve().parent / "benchmark_results.xlsx"
LARGE_INSTANCE_OUTPUT_FILE = (
    Path(__file__).resolve().parent / "benchmark_results_large_instances.xlsx"
)
INPUT_FILE_PATTERN = "*.pkl"
TIME_LIMIT_SECONDS = 180.0
START_CITY = 0
OVERWRITE_OUTPUT = True
LARGE_INSTANCE_MIN_CITY_COUNT = 16
LARGE_INSTANCE_RUNNABLE_ALGORITHMS = {"Nearest Neighbor", "ACO"}
RUN_LARGE_INSTANCES_ONLY = False

RESULT_HEADERS = (
    "data_file",
    "algorithm",
    "status",
    "path",
    "cost",
    "execution_time_ms",
    "execution_time_seconds",
    "peak_memory_bytes",
    "peak_memory_kb",
    "peak_memory_mb",
    "error",
)
TIME_EXCEEDED_STATUS = "time exceeded"
PATH_NOT_FOUND_STATUS = "path not found"
CITY_COUNT_PATTERN = re.compile(r"^tsp_data_(\d+)_")


def run_astar_admissible(
    cost_matrix: list[list[float]],
    total_cities: int,
    start_city: int = 0,
) -> Any:
    return run_astar(
        cost_matrix,
        total_cities,
        start_city,
        is_admissible=True,
    )


def run_astar_inadmissible(
    cost_matrix: list[list[float]],
    total_cities: int,
    start_city: int = 0,
) -> Any:
    return run_astar(
        cost_matrix,
        total_cities,
        start_city,
        is_admissible=False,
    )


ALGORITHMS: tuple[tuple[str, Callable[..., Any]], ...] = (
    ("BFS", run_bfs),
    ("DFS", run_dfs),
    ("Nearest Neighbor", run_nn),
    ("Dijkstra", run_dijkstra),
    ("A* admissible", run_astar_admissible),
    ("A* inadmissible", run_astar_inadmissible),
    ("ACO", run_aco),
)


def _measure_algorithm(
    result_queue: multiprocessing.Queue,
    algorithm: Callable[..., Any],
    cost_matrix: list[list[float]],
    total_cities: int,
    start_city: int,
) -> None:
    try:
        result = measure_performance(
            algorithm,
            cost_matrix,
            total_cities,
            start_city,
        )
        result_queue.put(("completed", result, ""))
    except Exception as error:
        result_queue.put(("error", None, f"{type(error).__name__}: {error}"))


def run_with_timeout(
    algorithm: Callable[..., Any],
    cost_matrix: list[list[float]],
    total_cities: int,
    start_city: int,
    timeout_seconds: float,
) -> tuple[str, tuple[Any, ...] | None, str]:
    context = multiprocessing.get_context("spawn")
    result_queue = context.Queue()
    process = context.Process(
        target=_measure_algorithm,
        args=(result_queue, algorithm, cost_matrix, total_cities, start_city),
    )
    process.start()
    process.join(timeout_seconds)

    if process.is_alive():
        process.terminate()
        process.join()
        result_queue.close()
        result_queue.join_thread()
        return TIME_EXCEEDED_STATUS, None, ""

    try:
        status, result, error = result_queue.get(timeout=1.0)
    except Empty:
        status = "error"
        result = None
        error = f"Worker exited with code {process.exitcode} without a result"
    finally:
        result_queue.close()
        result_queue.join_thread()

    return status, result, error


def load_input_file(path: Path) -> tuple[list[Any], list[list[float]]]:
    with path.open("rb") as file:
        cities, cost_matrix = pickle.load(file)
    return cities, cost_matrix


def create_workbook() -> Workbook:
    workbook = Workbook()
    worksheet = workbook.active
    if worksheet is None:
        raise RuntimeError("Workbook did not provide an active worksheet")
    worksheet.title = "Results"
    worksheet.append(RESULT_HEADERS)
    return workbook


def save_workbook_atomically(workbook: Workbook, output_file: Path) -> None:
    output_file.parent.mkdir(parents=True, exist_ok=True)
    temporary_file: str | None = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="wb",
            suffix=".xlsx",
            prefix=f".{output_file.stem}.",
            dir=output_file.parent,
            delete=False,
        ) as file:
            temporary_file = file.name
        workbook.save(temporary_file)
        os.replace(temporary_file, output_file)
    except Exception:
        if temporary_file is not None:
            temporary_path = Path(temporary_file)
            if temporary_path.exists():
                temporary_path.unlink()
        raise


def prepare_workbook(output_file: Path) -> Workbook:
    if output_file.exists() and not OVERWRITE_OUTPUT:
        return load_workbook(output_file)
    return create_workbook()


def input_file_sort_key(path: Path) -> tuple[int, str]:
    match = CITY_COUNT_PATTERN.match(path.stem)
    if match is None:
        return (2**63 - 1, path.name)
    return (int(match.group(1)), path.name)


def input_file_city_count(path: Path) -> int | None:
    match = CITY_COUNT_PATTERN.match(path.stem)
    return int(match.group(1)) if match is not None else None


def result_row(
    data_file: Path,
    algorithm_name: str,
    status: str,
    result: tuple[Any, ...] | None,
    error: str,
) -> tuple[Any, ...]:
    if result is None:
        return (
            data_file.name,
            algorithm_name,
            status,
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            error,
        )

    path, cost, execution_time_ms, peak_memory_bytes = result
    serialized_path = json.dumps(path) if path is not None else ""
    result_status = PATH_NOT_FOUND_STATUS if path is None else status
    return (
        data_file.name,
        algorithm_name,
        result_status,
        serialized_path,
        cost,
        execution_time_ms,
        execution_time_ms / 1000,
        peak_memory_bytes,
        peak_memory_bytes / 1024,
        peak_memory_bytes / (1024 * 1024),
        error,
    )


def _run_benchmark(
    output_file: Path,
    minimum_city_count: int | None = None,
    runnable_algorithms: set[str] | None = None,
) -> None:
    if TIME_LIMIT_SECONDS <= 0:
        raise ValueError("TIME_LIMIT_SECONDS must be greater than zero")
    if not DATA_DIRECTORY.is_dir():
        raise FileNotFoundError(f"Data directory does not exist: {DATA_DIRECTORY}")

    data_files = sorted(DATA_DIRECTORY.glob(INPUT_FILE_PATTERN), key=input_file_sort_key)
    if minimum_city_count is not None:
        data_files = [
            path
            for path in data_files
            if (city_count := input_file_city_count(path)) is not None
            and city_count >= minimum_city_count
        ]
    if not data_files:
        raise FileNotFoundError(
            f"No input files matching {INPUT_FILE_PATTERN!r} in {DATA_DIRECTORY}"
        )

    workbook = prepare_workbook(output_file)
    worksheet = workbook["Results"]

    for data_file in data_files:
        _, cost_matrix = load_input_file(data_file)
        total_cities = len(cost_matrix)

        for algorithm_name, algorithm in ALGORITHMS:
            if runnable_algorithms is None or algorithm_name in runnable_algorithms:
                status, result, error = run_with_timeout(
                    algorithm,
                    cost_matrix,
                    total_cities,
                    START_CITY,
                    TIME_LIMIT_SECONDS,
                )
            else:
                status, result, error = TIME_EXCEEDED_STATUS, None, ""
            row = result_row(data_file, algorithm_name, status, result, error)
            worksheet.append(row)
            save_workbook_atomically(workbook, output_file)
            print(f"{data_file.name} | {algorithm_name} | {status}")


def run_benchmark() -> None:
    _run_benchmark(OUTPUT_FILE)


def run_large_instances_benchmark() -> None:
    _run_benchmark(
        LARGE_INSTANCE_OUTPUT_FILE,
        minimum_city_count=LARGE_INSTANCE_MIN_CITY_COUNT,
        runnable_algorithms=LARGE_INSTANCE_RUNNABLE_ALGORITHMS,
    )


if __name__ == "__main__":
    multiprocessing.freeze_support()
    # run_benchmark()
    run_large_instances_benchmark()

