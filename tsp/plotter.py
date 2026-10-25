from __future__ import annotations

import math
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

import matplotlib.pyplot as plt


# Edit these settings before generating plots.
INPUT_FILE = Path(__file__).resolve().parent / "results_raw.xlsx"
OUTPUT_DIRECTORY = Path(__file__).resolve().parent / "plots"
OUTPUT_FORMAT = "png"
SHOW_PLOTS = False

ALGORITHM_ORDER = (
    "BFS",
    "DFS",
    "Nearest Neighbor",
    "Dijkstra",
    "A* admissible",
    "A* inadmissible",
    "ACO",
)

METRICS = (
    ("cost", "Route cost", "Cost"),
    ("execution_time_ms", "Execution time", "Time (ms)"),
    ("peak_memory_bytes", "Peak memory usage", "Memory (bytes)"),
)

DATA_FILE_PATTERN = re.compile(
    r"^tsp_data_(?P<city_count>\d+)_(?P<symmetry>asymmetric|symmetric)_"
    r"(?P<density>0_8|1_0)\.pkl$"
)


def load_results(input_file: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(input_file, read_only=True, data_only=True)
    worksheet = workbook.active
    if worksheet is None:
        raise RuntimeError(f"Workbook has no active worksheet: {input_file}")

    rows = worksheet.iter_rows(values_only=True)
    headers = next(rows, None)
    if headers is None:
        raise ValueError(f"Workbook is empty: {input_file}")

    header_indexes = {header: index for index, header in enumerate(headers)}
    required_headers = {
        "data_file",
        "algorithm",
        "status",
        "cost",
        "execution_time_ms",
        "peak_memory_bytes",
    }
    missing_headers = required_headers - header_indexes.keys()
    if missing_headers:
        raise ValueError(f"Missing columns in {input_file}: {sorted(missing_headers)}")

    results = []
    for row in rows:
        data_file = row[header_indexes["data_file"]]
        algorithm = row[header_indexes["algorithm"]]
        match = DATA_FILE_PATTERN.match(str(data_file))
        if match is None or algorithm not in ALGORITHM_ORDER:
            continue

        results.append(
            {
                "data_file": str(data_file),
                "algorithm": str(algorithm),
                "status": row[header_indexes["status"]],
                "city_count": int(match.group("city_count")),
                "scenario": (
                    match.group("symmetry"),
                    match.group("density"),
                ),
                "cost": row[header_indexes["cost"]],
                "execution_time_ms": row[
                    header_indexes["execution_time_ms"]
                ],
                "peak_memory_bytes": row[
                    header_indexes["peak_memory_bytes"]
                ],
            }
        )
    return results


def metric_value(result: dict[str, Any], metric_name: str) -> float | None:
    if result["status"] != "completed":
        return None

    value = result[metric_name]
    if value is None:
        return None

    numeric_value = float(value)
    return numeric_value if math.isfinite(numeric_value) else None


def scenario_label(scenario: tuple[str, str]) -> str:
    symmetry, density = scenario
    return f"{symmetry}, density {density.replace('_', '.')}"


def plot_scenario_metric(
    results: list[dict[str, Any]],
    scenario: tuple[str, str],
    metric_name: str,
    title: str,
    y_label: str,
) -> Path:
    figure, axis = plt.subplots(figsize=(12, 7))
    scenario_results = [
        result for result in results if result["scenario"] == scenario
    ]

    for algorithm in ALGORITHM_ORDER:
        algorithm_results = [
            result
            for result in scenario_results
            if result["algorithm"] == algorithm
        ]
        algorithm_results.sort(key=lambda result: result["city_count"])
        x_values = [result["city_count"] for result in algorithm_results]
        y_values = [
            metric_value(result, metric_name)
            for result in algorithm_results
        ]
        if x_values:
            axis.plot(
                x_values,
                y_values,
                marker="o",
                linestyle="--",
                linewidth=1.5,
                alpha=0.65,
                label=algorithm,
            )

    axis.set_title(f"{title} - {scenario_label(scenario)}")
    axis.set_xlabel("Number of cities (N)")
    axis.set_ylabel(y_label)
    axis.grid(True, linestyle="--", alpha=0.5)
    axis.legend()
    figure.tight_layout()

    output_name = (
        f"{scenario[0]}_{scenario[1]}_{metric_name}.{OUTPUT_FORMAT}"
    )
    output_path = OUTPUT_DIRECTORY / output_name
    figure.savefig(output_path, format=OUTPUT_FORMAT, dpi=150)
    if SHOW_PLOTS:
        plt.show()
    plt.close(figure)
    return output_path


def generate_plots() -> list[Path]:
    if not INPUT_FILE.is_file():
        raise FileNotFoundError(f"Results file does not exist: {INPUT_FILE}")

    OUTPUT_DIRECTORY.mkdir(parents=True, exist_ok=True)
    results = load_results(INPUT_FILE)
    scenarios = sorted({result["scenario"] for result in results})
    if len(scenarios) != 4:
        raise ValueError(
            f"Expected 4 scenarios, found {len(scenarios)}: {scenarios}"
        )

    generated_files = []
    for scenario in scenarios:
        for metric_name, title, y_label in METRICS:
            generated_files.append(
                plot_scenario_metric(
                    results,
                    scenario,
                    metric_name,
                    title,
                    y_label,
                )
            )
    return generated_files


if __name__ == "__main__":
    for path in generate_plots():
        print(f"Generated {path}")
